# SPDX-License-Identifier: MPL-2.0
"""The AESA export, exercised along the path the packaged app actually takes.

The export button did nothing on a single-product run, for two reasons that
TestClient-only tests could not see together: the UI returned early because no
DSM system was active, and the route 404'd because it required one. In the
packaged app a working route would still have produced nothing, because a
``blob:`` download is inert in that webview.

So these tests start a REAL uvicorn server, POST a serialised body over a real
socket with the same headers the client sends, and read the REAL file the
server wrote. Nothing is mocked between the request and the file on disk.
"""
from __future__ import annotations

import asyncio
import io
import json
import socket
import sys
import threading
import time
import urllib.request
from pathlib import Path

import pytest
import uvicorn
from openpyxl import load_workbook

from mapper.api.aesa import post_compute
from mapper.main import app
from mapper.models.aesa_schemas import AESAComputeRequest, ProspectiveSingleProductPoint
from tests.test_aesa_single_product import _single_product, _sp_config

YEARS = (2025, 2030, 2040)


@pytest.fixture(scope="module")
def server():
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        port = s.getsockname()[1]
    srv = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=port, log_level="warning"))
    t = threading.Thread(target=srv.run, daemon=True)
    t.start()
    for _ in range(200):
        if srv.started:
            break
        time.sleep(0.05)
    assert srv.started, "test server did not start"
    yield f"http://127.0.0.1:{port}/api"
    srv.should_exit = True
    t.join(timeout=10)


def _post(url: str, body: dict) -> tuple[int, bytes, dict]:
    # Exactly what the fixed client sends: a string body WITH the JSON header.
    req = urllib.request.Request(url, data=json.dumps(body).encode(), method="POST",
                                 headers={"Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            return r.status, r.read(), dict(r.headers)
    except urllib.error.HTTPError as e:
        return e.code, e.read(), {}


def _export_body() -> dict:
    """A real multi-year single-product AESA result, as the UI posts it: the
    config has NO DSM system, which is what used to make the route 404."""
    cfg = _sp_config()
    points = [ProspectiveSingleProductPoint(year=y, result=_single_product()) for y in YEARS]
    res = asyncio.run(post_compute(AESAComputeRequest(
        config=cfg, single_product_basis="prospective", prospective_single_product=points)))
    assert {r.year for r in res.results} == set(YEARS)
    return {"config": cfg.model_dump(), "result": res.model_dump(), "source_name": "EV9-NMC811-99.8kWh"}


@pytest.fixture
def home(tmp_path, monkeypatch):
    monkeypatch.setattr(Path, "home", classmethod(lambda cls: tmp_path))
    (tmp_path / "Downloads").mkdir()
    return tmp_path


def test_save_writes_a_real_file_to_downloads_and_says_where(server, home):
    status, raw, _ = _post(f"{server}/aesa/export/save", _export_body())
    assert status == 200, raw
    body = json.loads(raw)
    written = Path(body["path"])
    assert written.parent == home / "Downloads"
    assert written.is_file() and written.stat().st_size == body["bytes"] > 0
    assert "EV9-NMC811-99.8kWh" in written.name and written.name.endswith(".xlsx")


def test_the_saved_workbook_carries_configuration_and_every_year(server, home):
    status, raw, _ = _post(f"{server}/aesa/export/save", _export_body())
    assert status == 200, raw
    wb = load_workbook(Path(json.loads(raw)["path"]))

    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r and r[0]}
    assert summary["System"] == "EV9-NMC811-99.8kWh"
    assert summary["Impact basis (functional unit)"].startswith("per functional unit, as computed")
    assert summary["Allocated SOS basis"] == "per year"
    assert summary["Budget basis"] in ("CO2", "CO2e_GHG"), "the basis stays visible, read-only"
    assert "SSP scenario" in summary and "Carbon budget" in summary

    rows = list(wb["Impacts vs SOS"].iter_rows(values_only=True))
    head = rows[0]
    for col in ("Year", "PB ID", "Impact", "Allocated SOS", "SR", "Zone", "Unit",
                "Impact basis", "SOS basis"):
        assert col in head, col
    data = [dict(zip(head, r)) for r in rows[1:]]
    assert {d["Year"] for d in data} == set(YEARS), "every selected year is exported"
    for d in data:
        assert d["Zone"] in ("safe", "zone_of_uncertainty", "high_risk")
        assert d["SOS basis"] == "per year"
    assert "Carbon Budget" in wb.sheetnames


def test_the_download_route_works_without_a_dsm_system_too(server):
    status, raw, headers = _post(f"{server}/aesa/export", _export_body())
    assert status == 200, raw[:300]
    assert "spreadsheetml" in headers.get("content-type", headers.get("Content-Type", ""))
    assert load_workbook(io.BytesIO(raw))["Summary"]["B1"].value == "EV9-NMC811-99.8kWh"


@pytest.mark.skipif(sys.platform == "win32", reason="chmod cannot make a folder read-only on Windows")
def test_a_write_failure_is_reported_not_swallowed(server, home):
    (home / "Downloads").chmod(0o500)  # read-only
    try:
        status, raw, _ = _post(f"{server}/aesa/export/save", _export_body())
    finally:
        (home / "Downloads").chmod(0o700)
    assert status == 500 and b"Could not write" in raw
