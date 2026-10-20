# SPDX-License-Identifier: MPL-2.0
"""AESA sensitivity cases: selected, computed, returned and exported.

EVS40 run: "Sharing sensitivity" was ticked and the project has parameter
cases (v1g, v2h_credit0/50/100), yet no case could be seen by name and the
parameter cases never reached AESA at all -- AESA consumed one impact result.

Now the case window chooses which cases to compute: sharing principles
(computed by the existing sweep, narrowed to the selection) and parameter
cases (each recomputed through the single-product LCA handler). They come back
as ``result.cases`` and the export has one Sensitivity column group per case.

The parameter-case tests replace the LCA HANDLER only (CI has no ecoinvent):
everything from the request to the workbook on disk is the real path. The real
LCA under a case was checked against the EVS40 project in a running sidecar.
"""
from __future__ import annotations

import asyncio
import json
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from mapper.api import aesa as aesa_api
from mapper.api import lca as lca_api
from mapper.api import parameters as param_api
from mapper.api.dsm import _current_project
from mapper.models.aesa_schemas import AESAComputeRequest, ProspectiveSingleProductPoint
from mapper.models.parameter_schemas import Parameter, ParameterTable
from tests.test_aesa_export_real_path import _post, home, server  # noqa: F401  (fixtures)
from tests.test_aesa_single_product import _single_product, _sp_config


def _compute(**kw):
    return asyncio.run(aesa_api.post_compute(AESAComputeRequest(config=_sp_config(), **kw)))


def _climate(rows):
    return next(r for r in rows if r.pb_id == "climate_change")


# ── sharing cases ─────────────────────────────────────────────────────────────

def test_only_the_selected_sharing_principles_are_computed():
    res = _compute(single_product_result=_single_product(), run_sensitivity=True,
                   sharing_cases=["EpC", "AR"])
    assert [c.id for c in res.cases] == ["sharing:EpC", "sharing:AR"]
    assert set(res.sensitivity) == {"EpC", "AR"}, "the box plot sees the same selection"
    assert all(c.kind == "sharing" and c.definition for c in res.cases)


def test_an_empty_selection_computes_the_base_only():
    res = _compute(single_product_result=_single_product(), run_sensitivity=True, sharing_cases=[])
    assert res.cases == [] and not res.sensitivity


def test_no_selection_keeps_the_pre_window_behaviour():
    res = _compute(single_product_result=_single_product(), run_sensitivity=True)
    assert {c.id for c in res.cases} == {f"sharing:{p}" for p in res.sensitivity}
    assert len(res.cases) >= 3


# ── parameter cases ───────────────────────────────────────────────────────────

@pytest.fixture
def table():
    project = _current_project()
    param_api._tables[project] = ParameterTable(
        parameters={
            "ev_use": Parameter(name="ev_use", base_value=5820.5, scenario_overrides={"v1g": 5136.4}),
            "untouched": Parameter(name="untouched", base_value=1.0),
        },
        scenarios=["v1g"],
    )
    yield
    param_api._tables.pop(project, None)


@pytest.fixture
def lca_by_case(monkeypatch):
    """The LCA handler, replaced so a case changes the score: v1g halves it."""
    calls = []

    async def fake(body):
        calls.append(body)
        sp = _single_product()
        factor = 0.5 if body.parameter_scenario == "v1g" else 1.0
        for r in sp.results:
            r.score *= factor
        sp.parameter_scenario = body.parameter_scenario
        sp.compute_database = body.compute_database
        return sp

    monkeypatch.setattr(lca_api, "calculate_archetype_lca", fake)
    return calls


def test_a_parameter_case_is_recomputed_and_returned(table, lca_by_case):
    res = _compute(single_product_result=_single_product(), parameter_cases=["v1g"])
    (case,) = res.cases
    assert case.id == "parameter:v1g" and case.kind == "parameter"
    assert case.definition == ["ev_use = 5136.4 (base 5820.5)"], "only what differs"
    assert _climate(case.results).impact == pytest.approx(_climate(res.results).impact / 2)
    (call,) = lca_by_case
    assert call.parameter_scenario == "v1g" and call.archetype_id == "arc-bev"


def test_a_prospective_case_reruns_every_year(table, lca_by_case):
    points = [ProspectiveSingleProductPoint(year=y, result=_single_product()) for y in (2025, 2030)]
    res = _compute(single_product_basis="prospective", prospective_single_product=points,
                   parameter_cases=["v1g"])
    assert len(lca_by_case) == 2
    assert {r.year for r in res.cases[0].results} == {2025, 2030}


def test_an_unknown_case_is_refused(table, lca_by_case):
    with pytest.raises(HTTPException) as e:
        _compute(single_product_result=_single_product(), parameter_cases=["v1gg"])
    assert e.value.status_code == 400 and not lca_by_case


def test_a_fleet_source_refuses_parameter_cases(table):
    from mapper.core.aesa_engine import single_product_to_impact_result
    imp = single_product_to_impact_result(_single_product(), reference_year=2025)
    with pytest.raises(HTTPException) as e:
        _compute(impact_result=imp, parameter_cases=["v1g"])
    assert e.value.status_code == 400 and "single-product" in e.value.detail


# ── the export, through a real server, to a real file ─────────────────────────

def test_the_exported_sensitivity_sheet_has_one_column_group_per_selected_case(
    table, lca_by_case, server, home,  # noqa: F811
):
    res = _compute(single_product_result=_single_product(), run_sensitivity=True,
                   sharing_cases=["AR"], parameter_cases=["v1g"])
    status, raw, _ = _post(f"{server}/aesa/export/save", {
        "config": _sp_config().model_dump(), "result": res.model_dump(), "source_name": "EV9",
    })
    assert status == 200, raw
    wb = load_workbook(Path(json.loads(raw)["path"]))

    rows = list(wb["Sensitivity"].iter_rows(values_only=True))
    head = rows[0]
    assert head[:6] == ("Year", "PB ID", "Impact category", "Abbreviation", "Base SR", "Base zone")
    case_cols = [h for h in head[6:] if h.endswith(" SR")]
    assert len(case_cols) == 2
    assert case_cols[0].startswith("Sharing: ") and case_cols[1] == "Case: v1g SR"
    d = [dict(zip(head, r)) for r in rows[1:]]
    clim = next(x for x in d if x["PB ID"] == "climate_change")
    assert clim["Abbreviation"] == "GWP100"
    assert clim["Case: v1g ΔSR vs base"] == pytest.approx(clim["Case: v1g SR"] - clim["Base SR"])
    assert clim["Case: v1g SR"] == pytest.approx(clim["Base SR"] / 2)
    assert not any("EpC" in h for h in head), "only selected cases appear"

    summary = [r for r in wb["Summary"].iter_rows(values_only=True) if r and r[0]]
    defs = {r[0]: r[1] for r in summary}
    assert defs["Sensitivity"] == "2 case(s) beside the base"
    assert defs["Case: v1g"] == "ev_use = 5136.4 (base 5820.5)"
    assert "Sensitivity (long)" in wb.sheetnames
