# SPDX-License-Identifier: MPL-2.0
# This Source Code Form is subject to the terms of the Mozilla Public
# License, v. 2.0. If a copy of the MPL was not distributed with this
# file, You can obtain one at https://mozilla.org/MPL/2.0/.
#
# © Copyright 2026 Technical University of Denmark
# Lead developer: Leonardo Ferhati

"""Every workbook that reports impact or SR values states coverage for EVERY
indicator it reports -- never silently for some.

A workbook naming some indicators "not specified" and saying nothing about the
rest invites the reader to infer the rest were checked. So each indicator gets
one of three explicit statuses, and "not recorded" (a result computed before
the check existed) is never collapsed into "specified".
"""
from __future__ import annotations

import ast
from pathlib import Path

from openpyxl import Workbook

from mapper.core import coverage_export as ce
from mapper.models.authored_schemas import CoverageGap

BACKEND = Path(__file__).resolve().parents[1] / "mapper"
GW = ["EF v3.1", "climate change", "GWP100"]
AC = ["EF v3.1", "acidification", "accumulated exceedance (AE)"]
GAP = CoverageGap(method=AC, database="mine", code="abc", activity_name="Boiler", scope_note="CO2 only")


def _flat(ws) -> str:
    return "\n".join(" | ".join("" if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True))


def _wb(entries, **kw):
    wb = Workbook()
    wb.active.title = "Summary"
    wb.active.append(["Project", "P"])
    ce.finalize_coverage(wb, entries, **kw)
    return wb


# ── The three states ────────────────────────────────────────────────────────

def test_every_indicator_gets_a_status_including_the_specified_ones():
    wb = _wb([ce.CoverageEntry(None, [GAP], [GW, AC])])
    cov = _flat(wb["Coverage"])
    assert "EF v3.1 › climate change › GWP100 | specified" in cov, "a specified indicator must not be silent"
    assert "EF v3.1 › acidification › accumulated exceedance (AE) | NOT SPECIFIED | Boiler | mine | CO2 only" in cov
    assert "1 of 2 indicator(s) NOT SPECIFIED" in _flat(wb["Summary"])


def test_checked_and_clean_says_so():
    wb = _wb([ce.CoverageEntry(None, [], [GW, AC])])
    assert "Checked: all 2 indicator(s) specified" in _flat(wb["Summary"])
    assert "| NOT SPECIFIED |" not in _flat(wb["Coverage"])


def test_not_recorded_is_never_read_as_specified():
    wb = _wb([ce.CoverageEntry(None, None, [GW, AC])])
    cov = _flat(wb["Coverage"])
    assert "| specified |" not in cov
    assert cov.count("| not recorded |") == 2
    assert "NOT RECORDED for 2 of 2" in _flat(wb["Summary"])


def test_a_mixed_workbook_reports_both_unrecorded_and_not_specified():
    wb = _wb([ce.CoverageEntry("Old", None, [GW]), ce.CoverageEntry("New", [GAP], [GW, AC])],
             discriminator="Sensitivity case")
    line = _flat(wb["Summary"])
    assert "NOT RECORDED for 1 of 3" in line and "1 further indicator(s) NOT SPECIFIED" in line
    assert "Sensitivity case | Indicator | Status" in _flat(wb["Coverage"])


def test_nothing_on_the_first_sheet_moves():
    wb = _wb([ce.CoverageEntry(None, [], [GW])])
    assert wb["Summary"]["A1"].value == "Project" and wb["Summary"]["B1"].value == "P"
    assert wb.sheetnames == ["Summary", "Coverage"]


def test_aesa_rows_are_per_boundary_with_the_method_that_feeds_it():
    from mapper.models.aesa_schemas import AESACoverageGap

    g = AESACoverageGap(**GAP.model_dump(), pb_id="acidification")
    wb = _wb([ce.CoverageEntry(None, [g], [("climate_change", "Climate change"), ("acidification", "Acidification")],
                               aesa=True)])
    cov = _flat(wb["Coverage"])
    assert "Climate change |  | specified" in cov
    assert "Acidification | EF v3.1 › acidification › accumulated exceedance (AE) | NOT SPECIFIED | Boiler" in cov


# ── Real builders, all three states ─────────────────────────────────────────

def _static_result(gaps):
    from mapper.models.schemas import ArchetypeLCACalculateResult, ArchetypeLCAMethodResult

    kw = {} if gaps == "absent" else {"coverage_gaps": gaps}
    return ArchetypeLCACalculateResult(
        archetype_id="a", archetype_name="Heated product", scope="all", amount=1.0,
        stages_included=["Manufacturing"],
        results=[ArchetypeLCAMethodResult(method=m, method_label=m[-1], score=1.0, unit="u", contributions=[])
                 for m in (GW, AC)],
        **kw,
    )


def test_the_single_product_static_workbook_states_all_three():
    from mapper.api.impact import _build_single_product_static_workbook

    for gaps, expect in (([GAP], "1 of 2 indicator(s) NOT SPECIFIED"),
                         ([], "Checked: all 2"),
                         ("absent", "NOT RECORDED for 2 of 2")):
        wb = _build_single_product_static_workbook("Heated product", "all", [("Base", _static_result(gaps))])
        assert "Coverage" in wb.sheetnames
        assert expect in _flat(wb.worksheets[0]), (gaps, _flat(wb.worksheets[0])[-300:])


def test_an_old_stored_result_deserialises_as_not_recorded():
    """The default must be None: [] would claim 'checked, none'."""
    from mapper.models.schemas import ArchetypeLCACalculateResult

    old = _static_result("absent").model_dump(exclude={"coverage_gaps"})
    assert ArchetypeLCACalculateResult(**old).coverage_gaps is None


# ── Guard: no result workbook ships without the statement ──────────────────

#: Workbooks that report no impact or SR value, with the reason.
EXEMPT = {
    ("aesa.py", "_build_sharing_workbook"): "sharing-preset configuration; no results",
    ("bom.py", "_build_export_workbook"): "BOM rows (the authored inventory itself); no results",
    ("bom.py", "_build_multi_export_workbook"): "BOM rows; no results",
    ("dsm.py", "_build_export_workbook"): "DSM stock / flow counts; no impact values",
    ("parameters.py", "_build_workbook"): "the parameter table; no results",
    ("subsystems.py", "_build_subsystem_dsm_workbook"): "subsystem stock counts; no impact values",
}
CALLS = {"finalize_coverage", "_finalize_coverage"}


def _builders():
    for path in sorted((BACKEND / "api").glob("*.py")):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in tree.body:
            if isinstance(node, ast.FunctionDef) and node.name.startswith("_build") and node.name.endswith("workbook"):
                yield path.name, node


def test_every_result_workbook_states_coverage():
    missing = []
    for fname, fn in _builders():
        if (fname, fn.name) in EXEMPT:
            continue
        called = {n.func.id if isinstance(n.func, ast.Name) else getattr(n.func, "attr", "")
                  for n in ast.walk(fn) if isinstance(n, ast.Call)}
        if not called & CALLS:
            missing.append(f"{fname}:{fn.name}")
    assert missing == [], ("workbooks reporting results without a coverage statement "
                           "(add finalize_coverage, or EXEMPT with a reason): " + ", ".join(missing))


def test_every_exemption_still_exists():
    found = {(f, fn.name) for f, fn in _builders()}
    assert set(EXEMPT) <= found, f"stale exemptions: {set(EXEMPT) - found}"


def test_the_guard_sees_the_builders_it_should():
    names = {fn.name for _, fn in _builders()}
    assert {"_build_aesa_workbook", "_build_mfa_lca_workbook", "_build_multi_product_workbook",
            "_build_monte_carlo_workbook", "_build_contribution_workbook"} <= names
