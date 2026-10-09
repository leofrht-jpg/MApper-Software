# SPDX-License-Identifier: MPL-2.0
"""The Monte Carlo export when the impact-weighted coverage has no share.

``PedigreeCoverage.impact_share`` is ``float | None``. The Summary sheet used
to format it as ``impact_share * 100``, so posting a coverage whose share was
``None`` -- an archetype with nothing scoreable, or one whose scoreable rows
carry zero total impact -- returned a 500 ("unsupported operand type(s) for *:
'NoneType' and 'int"). Found walking 0.3.0 on a demo archetype showing
"0 of 2 materials scored".

Every case here goes through the REAL route with a serialised body, not the
builder: the crash was in the route's path, and a builder-only test is the
shape that let it ship.
"""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from mapper.main import app
from tests.test_monte_carlo_export import _coverage, _result

client = TestClient(app)


def _post(coverage: dict | None):
    body = {"result": _result().model_dump()}
    if coverage is not None:
        body["coverage"] = coverage
    return client.post("/api/lca/monte-carlo/export", json=body)


def _summary(resp) -> dict:
    wb = load_workbook(io.BytesIO(resp.content))
    return {
        r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True)
        if r and r[0]
    }


def _cov(**over) -> dict:
    d = _coverage().model_dump()
    d.update(over)
    return d


@pytest.mark.parametrize(
    "case",
    [
        # scoreable rows, zero total impact -> no share (the reported repro:
        # "0 of 2 materials scored")
        dict(archetype_materials_total=2, archetype_materials_scored=0, top_unscored=[]),
        # nothing scoreable: every row is a parameter expression
        dict(archetype_materials_total=0, archetype_materials_scored=0, top_unscored=[]),
    ],
    ids=["zero-impact", "nothing-scoreable"],
)
def test_coverage_present_with_no_share_exports(case):
    """The cross: coverage IS supplied AND its share is None."""
    r = _post(_cov(impact_share=None, **case))
    assert r.status_code == 200, r.text
    text = str(_summary(r)["Impact-weighted coverage"])
    # Never formatted as a number, never the string "None", never a fake 0%.
    assert "None" not in text
    assert not text.startswith("0.0%")
    assert "climate change | GWP100" in text


def test_zero_impact_case_says_why_and_how_many_are_scored():
    r = _post(_cov(impact_share=None, archetype_materials_total=2,
                   archetype_materials_scored=0, top_unscored=[]))
    text = _summary(r)["Impact-weighted coverage"]
    assert "zero total" in text and "0 of 2 scored" in text


def test_nothing_scoreable_case_points_at_the_parameters():
    r = _post(_cov(impact_share=None, archetype_materials_total=0,
                   archetype_materials_scored=0, top_unscored=[]))
    text = _summary(r)["Impact-weighted coverage"]
    assert "parameter expression" in text


def test_a_real_share_still_formats_as_a_percentage():
    r = _post(_cov(impact_share=0.82))
    assert r.status_code == 200
    assert _summary(r)["Impact-weighted coverage"] == "82.0% of climate change | GWP100"


def test_absent_coverage_is_still_stated():
    r = _post(None)
    assert r.status_code == 200
    assert "not recorded" in _summary(r)["Impact-weighted coverage"]
