# SPDX-License-Identifier: MPL-2.0
# This Source Code Form is subject to the terms of the Mozilla Public
# License, v. 2.0. If a copy of the MPL was not distributed with this
# file, You can obtain one at https://mozilla.org/MPL/2.0/.
#
# © Copyright 2026 Technical University of Denmark
# Lead developer: Leonardo Ferhati

"""Paired multi-item Monte Carlo over items in DIFFERENT databases.

``bw2calc.MonteCarloLCA`` loads its matrices from the databases reachable from
the demand it is CONSTRUCTED with. The paired run constructed it from item 0's
demand only, so a second item living in a database item 0 never reaches was
outside the technosphere: ``OutsideTechnosphere`` on the first iteration. The
0.3.0 worked examples hit it directly -- the polycarbonate enclosure ("Supplier
PCFs (worked example)") paired with the flare ("Flaring (worked example)") --
because two authored databases share nothing but biosphere.

Nothing in CI's ecoinvent-free environment could reach this before: every
multi-item test either mocked the run or used items over one database. So this
builds a real (tiny) Brightway world -- one biosphere, one method, two
single-activity databases -- in a TEMPORARY project, deleted afterwards, and
drives the real route.

The world lives in a THROWAWAY Brightway base directory
(``projects._use_temp_directory``), never in the developer's real one, and the
directory is removed tolerantly. On Windows a plain ``delete_project(...,
delete_dir=True)`` failed with WinError 32 on ``lci/databases.db``: bw2data's
peewee ``SqliteDatabase`` keeps one connection PER THREAD, and switching
projects closes only the calling thread's. The Monte Carlo worker and the
TestClient threads had their own open handles, and Windows will not unlink an
open file. A leftover temp file is not a test failure.
"""
from __future__ import annotations

import gc
import math
import shutil
import time
import uuid

import pytest
from fastapi.testclient import TestClient

bw2data = pytest.importorskip("bw2data")
pytest.importorskip("bw2calc")

BIO = "pt-bio"
METHOD = ("PT", "climate", "gwp")


def _activity_db(name: str, kg_co2: float) -> dict:
    """One process, one lognormal CO2 exchange -- the shape of an authored db."""
    return {
        (name, "a1"): {
            "name": f"{name} activity", "unit": "kilogram", "location": "GLO",
            "type": "process", "reference product": f"{name} product",
            "exchanges": [
                {"input": (name, "a1"), "amount": 1.0, "type": "production"},
                {"input": (BIO, "co2"), "amount": kg_co2, "type": "biosphere",
                 "uncertainty type": 2, "loc": math.log(kg_co2), "scale": 0.1},
            ],
        }
    }


@pytest.fixture()
def two_databases():
    from mapper.api import bom
    from mapper.models.bom_schemas import Archetype, BOMNode, EcoinventLink

    projects = bw2data.projects
    original = projects.current
    project = f"pytest-mc-two-dbs-{uuid.uuid4().hex[:8]}"
    temp_dir = projects._use_temp_directory()      # throwaway base dir, not the real one
    saved = None
    try:
        projects.set_current(project)
        bw2data.Database(BIO).write({
            (BIO, "co2"): {"name": "Carbon dioxide, fossil", "unit": "kilogram",
                           "type": "emission", "categories": ("air",)},
        })
        bw2data.Method(METHOD).register(unit="kg CO2-eq")
        bw2data.Method(METHOD).write([((BIO, "co2"), 1.0)])
        bw2data.Database("auth-enclosure").write(_activity_db("auth-enclosure", 2.8))
        bw2data.Database("auth-flare").write(_activity_db("auth-flare", 3.3))

        def arc(aid, db):
            return Archetype(id=aid, name=aid, bom=[BOMNode(
                name="Manufacturing", node_type="component", quantity=1.0, unit="unit",
                children=[BOMNode(name=f"{db} row", node_type="material", quantity=1.0,
                                  unit="kilogram",
                                  ecoinvent_activity=EcoinventLink(database=db, code="a1",
                                                                   name=f"{db} activity"))])])

        saved = bom._archetypes.get(project)
        bom._archetypes[project] = {"enclosure": arc("enclosure", "auth-enclosure"),
                                    "flare": arc("flare", "auth-flare")}
        yield project
    finally:
        bom._archetypes.pop(project, None)
        if saved is not None:
            bom._archetypes[project] = saved
        # Back to the real base dir and the project the test started on; this
        # closes this thread's handles on the temp sqlite files.
        projects._restore_orig_directory()
        projects.set_current(original, update=False)
        # Handles held by finished worker threads go with their thread-locals.
        gc.collect()
        shutil.rmtree(temp_dir, ignore_errors=True)


def _run(ids, n=12, seed=7):
    from mapper.main import app

    c = TestClient(app)
    r = c.post("/api/lca/monte-carlo/multi", json={
        "archetype_ids": ids, "methods": [list(METHOD)], "scope": "all",
        "iterations": n, "seed": seed, "keep_samples": True})
    assert r.status_code == 200, r.text
    tid = r.json()["task_id"]
    for _ in range(300):
        g = c.get(f"/api/lca/monte-carlo/multi/{tid}")
        if g.status_code != 409:
            return c, g
        time.sleep(0.1)
    raise AssertionError("paired run did not finish")


def test_items_in_two_unrelated_databases_run_paired(two_databases):
    _, g = _run(["enclosure", "flare"])
    assert g.status_code == 200, g.text   # was 500: OutsideTechnosphere
    items = {it["archetype_name"]: it for it in g.json()["items"]}
    assert len(items) == 2
    medians = sorted(d["median"] for it in items.values() for d in it["distributions"])
    # Each solved against its OWN activity, not the other's (2.8 vs 3.3 kg).
    assert medians[0] == pytest.approx(2.8, rel=0.15)
    assert medians[1] == pytest.approx(3.3, rel=0.15)


def test_the_pair_exports(two_databases):
    c, g = _run(["enclosure", "flare"])
    assert g.status_code == 200, g.text
    x = c.post("/api/lca/monte-carlo/multi/export", json={"result": g.json()})
    assert x.status_code == 200, x.text
    from io import BytesIO

    from openpyxl import load_workbook

    assert "Distributions" in load_workbook(BytesIO(x.content)).sheetnames


def test_the_union_leaves_a_one_background_run_untouched(two_databases):
    """The ordinary case -- every item over the same background -- must not
    move a number. The helper therefore leaves item 0's database list exactly
    as bw2calc built it when the union adds nothing, and only APPENDS when it
    does (so item 0's databases keep their positions either way)."""
    import bw2calc

    from mapper.api import monte_carlo as MC

    enclosure = {("auth-enclosure", "a1"): 1.0}
    flare = {("auth-flare", "a1"): 1.0}

    mc = bw2calc.MonteCarloLCA(enclosure, METHOD, seed=1)
    before = list(mc.database_filepath)
    MC._union_database_filepaths(mc, [enclosure, {("auth-enclosure", "a1"): 3.0}])
    assert mc.database_filepath == before          # same list, same order

    MC._union_database_filepaths(mc, [enclosure, flare])
    assert mc.database_filepath[: len(before)] == before
    assert len(mc.database_filepath) == len(before) + 1
    assert "auth-flare" in mc.database_filepath[-1]
