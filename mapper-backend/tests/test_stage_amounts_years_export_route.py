# SPDX-License-Identifier: MPL-2.0
"""The Lifetime preset is gone; the export's "Lifetime (yr)" row still behaves.

Through the real route with the request a browser sends (raw JSON, the header
set), reading the workbook the server returns. Three runs:

* Custom with years entered (the only way a run now carries a lifetime): the
  row prints those years;
* Custom typed stage by stage, and "1 year": "not used", never a stale number;
* a result made BEFORE the preset was removed (``preset="lifetime"``): it still
  exports with its lifetime, so no saved result breaks.
"""
from __future__ import annotations

import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

RESULT = {
    "archetype_id": "ev9", "archetype_name": "EV9", "scope": "all", "amount": 1.0,
    "stage_amounts": {"Manufacturing": 1, "Use Phase": 16}, "stages_included": ["Manufacturing", "Use Phase"],
    "results": [{"method": ["EF v3.1", "climate change", "GWP100"], "method_label": "GWP100",
                 "score": 3249.6, "unit": "kg CO2-eq", "contributions": []}],
    "elapsed_seconds": 0.1, "warnings": [], "stage_breakdown": None,
}


def _export(meta: dict) -> dict:
    from mapper.main import app

    body = {"archetype_name": "EV9", "scope": "all",
            "scenarios": [{"label": "Base", "result": RESULT}], "stage_amounts_meta": meta}
    r = TestClient(app).post("/api/impact/export-single-product-static", content=json.dumps(body),
                             headers={"Content-Type": "application/json"})
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content))["Configuration"]
    return {row[0].value: row[1].value for row in ws.iter_rows(max_col=2) if row[0].value is not None}


@pytest.mark.parametrize(("meta", "preset", "lifetime"), [
    ({"preset": "custom", "lifetime": 16, "amounts": {"Manufacturing": 1, "Use Phase": 16}}, "custom", 16),
    ({"preset": "custom", "lifetime": None, "amounts": {"Manufacturing": 1, "Use Phase": 16}},
     "custom", "not used (custom stage amounts, no years entered)"),
    ({"preset": "1year", "lifetime": None, "amounts": {"Manufacturing": 1, "Use Phase": 1}},
     "1year", "not used (1-year preset)"),
    ({"preset": "lifetime", "lifetime": 16, "amounts": {"Manufacturing": 1, "Use Phase": 16}}, "lifetime", 16),
], ids=["custom-with-years", "custom-typed", "1year", "legacy-lifetime-result"])
def test_the_lifetime_row(meta, preset, lifetime):
    cfg = _export(meta)
    assert cfg["Preset"] == preset
    assert cfg["Lifetime (yr)"] == lifetime


def test_a_meta_without_a_lifetime_field_still_exports():
    """The new frontend may omit it; the schema default is None, not 15 or 1."""
    cfg = _export({"preset": "custom", "amounts": {"Use Phase": 16}})
    assert cfg["Lifetime (yr)"].startswith("not used")
