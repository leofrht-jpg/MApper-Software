# SPDX-License-Identifier: MPL-2.0
# This Source Code Form is subject to the terms of the Mozilla Public
# License, v. 2.0. If a copy of the MPL was not distributed with this
# file, You can obtain one at https://mozilla.org/MPL/2.0/.
#
# © Copyright 2026 Technical University of Denmark
# Lead developer: Leonardo Ferhati

"""Where an authored exchange's uncertainty comes from.

``ecoinvent`` (the default, and what every exchange written before this was):
the amount is an emission, so ecoinvent's spread for that flow is the reference
-- its median is the basic variance and the GSD2 floor applies.

``supplied``: the amount is a figure whose uncertainty belongs to whoever
produced it (a supplier PCF, an EPD). It is a characterised RESULT, so
ecoinvent's per-flow median describes a different quantity: the author states
the variance and its source, and no floor is checked.

The exemption must stay hard to reach. The floor exists because authored
inventories understate, so ``supplied`` asks for MORE than the default (a
variance AND its source), never less, and is never the default.
"""
from __future__ import annotations

import pytest

from mapper.core import authored_engine as eng
from mapper.models.authored_schemas import ExchangeInput

from tests.test_authored_databases import CO2, NOX, PED_HIGH, PED_LOW, FakeBackend

# NOX in the fixture carries real MAp-test-shaped stats: median bv 0.08,
# floor GSD2 1.93 over 2422 lognormal exchanges.
SUPPLIED = {"basic_variance": 0.0083, "basic_variance_reason": "supplier PCF: no uncertainty reported"}


def _ex(**kw) -> ExchangeInput:
    return ExchangeInput(flow_database=NOX.database, flow_code=NOX.code, amount=1.0,
                         pedigree=kw.pop("pedigree", PED_LOW), **kw)


def _resolve(**kw):
    return eng.resolve_exchange(_ex(**kw), FakeBackend())


# ── The default ─────────────────────────────────────────────────────────────

def test_the_default_is_ecoinvent_and_unchanged():
    assert ExchangeInput(flow_database="biosphere3", flow_code="x", amount=1.0,
                         pedigree=PED_HIGH).uncertainty_basis == "ecoinvent"
    e = _resolve(pedigree=PED_HIGH)
    assert e.uncertainty_basis == "ecoinvent"
    assert e.basic_variance_source == "ecoinvent_median" and e.basic_variance == 0.08
    assert e.floor_status == "floored" and e.floor_gsd2 == 1.93


def test_the_ecoinvent_basis_still_refuses_a_user_variance_and_a_low_gsd2():
    """Unchanged: this is the protection the exemption must not erode."""
    with pytest.raises(eng.AuthoredError) as e:
        _resolve(pedigree=PED_HIGH, **SUPPLIED)
    assert e.value.codes == ["bv_not_settable"]
    with pytest.raises(eng.AuthoredError) as e:
        _resolve(pedigree=PED_LOW)
    assert e.value.codes == ["below_floor"]


# ── The supplied basis ──────────────────────────────────────────────────────

def test_supplied_uses_the_authors_variance_and_checks_no_floor():
    e = _resolve(uncertainty_basis="supplied", **SUPPLIED)
    assert e.uncertainty_basis == "supplied"
    assert e.basic_variance == 0.0083, "ecoinvent's median must not be substituted"
    assert e.basic_variance_source == "supplied"
    assert e.basic_variance_detail == SUPPLIED["basic_variance_reason"]
    # PED_LOW on the ecoinvent basis is refused as below_floor (1.93). Here the
    # floor is not a bound at all: not applied, and not silently passed either.
    assert e.floor_status == "not_applicable" and e.floor_gsd2 is None
    assert e.floor_reason is None
    assert "does not apply" in e.floor_detail and "characterised result" in e.floor_detail
    assert e.gsd2 < 1.93


def test_supplied_is_more_work_than_the_default_never_less():
    """The input the ecoinvent basis ACCEPTS is refused on the supplied one."""
    assert _resolve(pedigree=PED_HIGH).basic_variance == 0.08          # accepted as-is
    with pytest.raises(eng.AuthoredError) as e:
        _resolve(pedigree=PED_HIGH, uncertainty_basis="supplied")      # same input
    assert e.value.codes == ["bv_required", "bv_reason_required"]


def test_supplied_requires_a_source_for_the_variance():
    with pytest.raises(eng.AuthoredError) as e:
        _resolve(uncertainty_basis="supplied", basic_variance=0.0083)
    assert e.value.codes == ["bv_reason_required"]
    assert "whose uncertainty this is" in e.value.problems[0]


def test_supplied_refuses_a_floor_reason_rather_than_storing_a_meaningless_one():
    with pytest.raises(eng.AuthoredError) as e:
        _resolve(uncertainty_basis="supplied", floor_reason="tighter than ecoinvent", **SUPPLIED)
    assert e.value.codes == ["floor_reason_not_applicable"]


def test_supplied_is_available_where_ecoinvent_has_no_statistics_too():
    """The basis is about what the number IS, not about what ecoinvent holds."""
    ex = ExchangeInput(flow_database=CO2.database, flow_code=CO2.code, amount=3.5,
                       pedigree=PED_HIGH, uncertainty_basis="supplied", **SUPPLIED)
    e = eng.resolve_exchange(ex, FakeBackend())
    assert e.floor_status == "not_applicable" and e.basic_variance == 0.0083


def test_a_supplied_exchange_is_not_reported_as_unfloored():
    """'unfloored' means the check COULD NOT run; 'not_applicable' means it does
    not apply. An activity must not claim the first for the second."""
    from mapper.models.authored_schemas import ActivityInput

    act = eng.build_activity(ActivityInput(
        name="supplier material", unit="kilogram", scope="partial",
        scope_note="supplier PCF only",
        exchanges=[_ex(uncertainty_basis="supplied", **SUPPLIED)]), FakeBackend(), eng.new_code())
    assert act.exchanges[0].floor_status == "not_applicable"
    assert act.has_unfloored_exchange is False


# ── Migration ───────────────────────────────────────────────────────────────

def test_an_exchange_written_before_the_field_is_ecoinvent():
    stored = _resolve(pedigree=PED_HIGH)
    legacy = stored.model_dump(exclude={"uncertainty_basis"})
    from mapper.models.authored_schemas import AuthoredExchange

    assert AuthoredExchange(**legacy).uncertainty_basis == "ecoinvent"


def test_the_basis_is_outside_the_fingerprint():
    """Adding the field must not mark existing authored databases stale."""
    from mapper.models.authored_schemas import AuthoredDatabase, AuthoredActivity

    def db(basis):
        e = _resolve(pedigree=PED_HIGH) if basis == "ecoinvent" else _resolve(
            uncertainty_basis="supplied", **SUPPLIED)
        a = AuthoredActivity(code=eng.new_code(), name="a", reference_product="a", unit="kg",
                             location="GLO", scope="complete", exchanges=[e])
        return AuthoredDatabase(name="mine", created_at="t", updated_at="t", activities=[a])

    same = db("ecoinvent")
    assert eng.fingerprint(same) == eng.fingerprint(same.model_copy(deep=True))
    # the basis alone does not move it; the variance it implies does, which is
    # correct -- that number IS written to Brightway.
    flipped = same.model_copy(deep=True)
    flipped.activities[0].exchanges[0].uncertainty_basis = "supplied"
    assert eng.fingerprint(flipped) == eng.fingerprint(same)


def test_the_basis_never_reaches_brightway():
    from mapper.models.authored_schemas import AuthoredActivity, AuthoredDatabase

    e = _resolve(uncertainty_basis="supplied", **SUPPLIED)
    a = AuthoredActivity(code=eng.new_code(), name="a", reference_product="a", unit="kg",
                         location="GLO", scope="complete", exchanges=[e])
    written = str(eng.to_bw2(AuthoredDatabase(name="mine", created_at="t", updated_at="t",
                                              activities=[a])))
    assert "uncertainty_basis" not in written and "supplied" not in written
    # the sigma it produced IS written -- that is the number Brightway samples
    assert str(e.sigma)[:8] in written


# ── Into the result, not just the editor ────────────────────────────────────

def _defs_with(*exchanges):
    from mapper.models.authored_schemas import AuthoredActivity, AuthoredDatabase, AuthoredDefinitions

    a = AuthoredActivity(code=eng.new_code(), name="supplier material", reference_product="m",
                         unit="kilogram", location="GLO", scope="complete", exchanges=list(exchanges))
    return AuthoredDefinitions(databases=[AuthoredDatabase(
        name="mine", created_at="t", updated_at="t", activities=[a])]), a


def test_the_notes_name_both_bases_never_only_the_exception(monkeypatch):
    """A block listing only 'supplied' would leave the reader to assume the
    rest are ecoinvent-referenced -- the same silence the coverage statement
    exists to avoid."""
    from mapper.core import authored_coverage as cov
    from mapper.core import authored_storage

    supplied = eng.resolve_exchange(ExchangeInput(
        flow_database=CO2.database, flow_code=CO2.code, amount=3.5, pedigree=PED_LOW,
        uncertainty_basis="supplied", **SUPPLIED), FakeBackend())
    defs, act = _defs_with(_resolve(pedigree=PED_HIGH), supplied)
    authored_storage.save_definitions("P", defs)
    notes = cov.uncertainty_notes({("mine", act.code)}, "P")
    assert [n.basis for n in notes] == ["ecoinvent", "supplied"]
    assert notes[0].detail.startswith("floored:")
    assert "does not apply" in notes[1].detail
    assert notes[1].gsd2 == pytest.approx(
        eng.gsd2_from_sigma(eng.total_sigma(PED_LOW, 0.0083)))


def test_an_unused_authored_activity_is_not_listed(monkeypatch):
    from mapper.core import authored_coverage as cov
    from mapper.core import authored_storage

    defs, _act = _defs_with(_resolve(pedigree=PED_HIGH))
    authored_storage.save_definitions("P", defs)
    assert cov.uncertainty_notes({("ecoinvent-3.10-cutoff", "x")}, "P") == []


def _note(basis="supplied", **kw):
    from mapper.models.authored_schemas import AuthoredUncertainty

    return AuthoredUncertainty(database="mine", code="c" * 32, activity_name="supplier material",
                               flow_name="Carbon dioxide, fossil", flow_categories=["air"],
                               basic_variance=0.0083, gsd2=1.229, basis=basis,
                               detail=eng.SUPPLIED_DETAIL if basis == "supplied" else "floored: …", **kw)


def test_the_coverage_sheet_carries_the_basis_block():
    from openpyxl import Workbook

    from mapper.core import coverage_export as ce

    wb = Workbook()
    wb.active.title = "Summary"
    wb.active.append(["Project", "P"])
    ce.finalize_coverage(wb, [ce.CoverageEntry(None, [], [["EF v3.1", "climate change", "GWP100"]])],
                         notes=[_note(), _note(basis="ecoinvent")])
    txt = "\n".join(" | ".join("" if v is None else str(v) for v in r)
                    for r in wb["Coverage"].iter_rows(values_only=True))
    assert "Authored exchanges: where each uncertainty comes from" in txt
    assert "supplier material | mine | Carbon dioxide, fossil [air] | 0.0083 | 1.229 | supplied" in txt
    assert "| ecoinvent |" in txt
    assert "do not compare a supplied gsd2" in txt.lower()


def test_the_monte_carlo_pedigree_sheet_carries_it_too():
    """Authored exchanges are sampled as background, so they are absent from the
    scored-input rows -- and this is the sheet where GSD2s get compared."""
    from tests.test_monte_carlo_export import _coverage, _result
    from mapper.api.monte_carlo import _build_monte_carlo_workbook

    wb = _build_monte_carlo_workbook(_result(authored_uncertainty=[_note()]), _coverage())
    txt = "\n".join(" | ".join("" if v is None else str(v) for v in r)
                    for r in wb["Pedigree scores"].iter_rows(values_only=True))
    assert "Authored exchanges: where each uncertainty comes from" in txt
    assert "supplied" in txt and "does not apply" in txt


def test_the_paired_workbook_lists_each_exchange_once():
    from tests.test_paired_multi_export import _result
    from mapper.api.monte_carlo import _build_monte_carlo_multi_workbook

    res = _result()
    res.items[0] = res.items[0].model_copy(update={"authored_uncertainty": [_note()]})
    res.items[1] = res.items[1].model_copy(update={"authored_uncertainty": [_note()]})
    txt = "\n".join(" | ".join("" if v is None else str(v) for v in r)
                    for r in _build_monte_carlo_multi_workbook(res)["Summary"].iter_rows(values_only=True))
    assert txt.count("supplier material | mine") == 1
